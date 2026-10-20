"""Write the results somewhere people can use them.

- Excel report: one sheet per question, with charts, for people who live in Excel
- Power BI export: one clean "fact table" CSV per run, ready to load into Power BI
- PNG chart of the daily cost with the spikes marked (used in the README)
"""

from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analysis import usage_rows

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0.00'
PERCENT = '0.0%'

POWERBI_COLUMNS = {
    "Date": "Date",
    "Month": "Month",
    "ProviderName": "Provider",
    "ServiceCategory": "ServiceCategory",
    "ServiceName": "Service",
    "SubAccountName": "Account",
    "RegionName": "Region",
    "tag_business_unit": "Team",
    "tag_environment": "Environment",
    "tag_application": "Application",
    "ChargeCategory": "ChargeCategory",
    "IsTagged": "IsTagged",
}


def write_excel(results: dict, path) -> Path:
    """Excel workbook with a summary sheet, breakdowns, spikes and recommendations."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cost_col = results["cost_column"]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _summary_sheet(results).to_excel(writer, sheet_name="Summary", index=False)

        tips = pd.DataFrame({"#": range(1, len(results["recommendations"]) + 1),
                             "Recommendation": results["recommendations"]})
        tips.to_excel(writer, sheet_name="Recommendations", index=False)

        sheets = {
            "By service": results["by_service"],
            "By team": results["by_team"],
            "By account": results["by_account"],
            "By environment": results["by_environment"],
            "By provider": results["by_provider"],
            "By region": results["by_region"],
            "Top resources": results["top_resources"],
            "Untagged": results["untagged_by_account"],
        }
        for name, table in sheets.items():
            table.to_excel(writer, sheet_name=name, index=False)

        daily = results["daily"].rename_axis("Date").reset_index(name="Cost")
        daily.to_excel(writer, sheet_name="Daily", index=False)
        results["anomalies"].to_excel(writer, sheet_name="Spikes", index=False)

        book = writer.book
        for sheet in book.worksheets:
            _style(sheet)
        _bar_chart(book["By service"], f"Usage cost by service ({cost_col})", "E2")
        _bar_chart(book["By team"], f"Usage cost by team ({cost_col})", "E2")
        _line_chart(book["Daily"], "Daily usage cost", "D2")
        book["Recommendations"].column_dimensions["B"].width = 110
        for row in book["Recommendations"].iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")

    return path


def write_powerbi(df: pd.DataFrame, path, cost: str = "BilledCost") -> Path:
    """One row per day x service x account x team: small, tidy, easy to model in Power BI."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    keep = list(POWERBI_COLUMNS)
    costs = ["BilledCost", "EffectiveCost", "ListCost"]
    fact = (
        df[keep + costs]
        .groupby(keep, dropna=False)[costs].sum()
        .reset_index()
        .rename(columns=POWERBI_COLUMNS)
    )
    fact["Date"] = pd.to_datetime(fact["Date"]).dt.strftime("%Y-%m-%d")
    fact.to_csv(path, index=False)
    return path


def write_chart(results: dict, path) -> Path:
    """PNG of the daily cost with spike days marked in red."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    daily = results["daily"]
    spikes = results["anomalies"]

    fig, ax = plt.subplots(figsize=(10, 4))
    ax.plot(daily.index, daily.values, color="#1F3A5F", linewidth=2, label="daily usage cost")
    if not spikes.empty:
        ax.scatter(spikes["Date"], spikes["Cost"], color="#D1495B", s=80, zorder=3, label="spike")
    ax.set_title(f"Daily cloud usage cost ({results['month']})")
    ax.set_ylabel(results["cost_column"])
    ax.grid(alpha=0.3)
    ax.legend()
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(path, dpi=120)
    plt.close(fig)
    return path


def console_summary(results: dict) -> str:
    """Short text version for the terminal."""
    s = results["summary"]
    q = results["quality"]
    lines = [
        f"period          {s['period_start']} to {s['period_end']}  (month: {results['month']})",
        f"providers       {', '.join(s['providers'])}  |  {s['accounts']} accounts  |  {s['resources']:,} resources",
        f"usage cost      {s['usage_cost']:>12,.2f}",
        f"credits         {s['credits']:>12,.2f}",
        f"net cost        {s['net_cost']:>12,.2f}",
        f"untagged        {s['untagged_cost']:>12,.2f}  ({s['untagged_share']:.0%} of usage)",
        f"cleaned data    {q.get('rows_in', 0):,} rows in, {q.get('category_fixed', 0)} categories fixed, "
        f"{q.get('duplicates_removed', 0)} duplicates, {results['rows_outside_month']} rows outside the month",
        "",
        "top services:",
    ]
    for _, row in results["by_service"].head(5).iterrows():
        lines.append(f"  {row.iloc[0]:<40} {row['Cost']:>10,.2f}  {row['Share']:>5.0%}")
    lines.append("")
    lines.append("recommendations:")
    for i, tip in enumerate(results["recommendations"], 1):
        lines.append(f"  {i}. {tip}")
    return "\n".join(lines)


def _summary_sheet(results: dict) -> pd.DataFrame:
    s = results["summary"]
    rows = [
        ("Period", f"{s['period_start']} to {s['period_end']}"),
        ("Cost column", results["cost_column"]),
        ("Providers", ", ".join(s["providers"])),
        ("Accounts", s["accounts"]),
        ("Resources", s["resources"]),
        ("Usage cost", round(s["usage_cost"], 2)),
        ("Credits", round(s["credits"], 2)),
        ("Net cost (after credits)", round(s["net_cost"], 2)),
        ("List price of usage", round(s["list_cost"], 2)),
        ("Saved by discounts", round(s["discount_savings"], 2)),
        ("Untagged usage cost", round(s["untagged_cost"], 2)),
        ("Untagged share", f"{s['untagged_share']:.1%}"),
        ("Non-prod cost on weekends", round(results["nonprod_weekend"]["weekend_cost"], 2)),
        ("Compute covered by commitments", f"{results['commitment']['coverage']:.1%}"),
        ("Cost spikes found", len(results["anomalies"])),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def _style(sheet) -> None:
    """Header colours, number formats, column widths."""
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    sheet.freeze_panes = "A2"

    headers = [c.value for c in sheet[1]]
    for idx, header in enumerate(headers, 1):
        letter = get_column_letter(idx)
        values = [str(c.value) for c in sheet[letter] if c.value is not None]
        sheet.column_dimensions[letter].width = min(max(len(v) for v in values) + 2, 60) if values else 12
        fmt = None
        if header in ("Cost", "Expected", "Extra", "Untagged cost"):
            fmt = MONEY
        elif header in ("Share", "Increase"):
            fmt = PERCENT
        elif header == "Date":
            fmt = "DD.MM.YYYY"
        if fmt:
            for cell in sheet[letter][1:]:
                cell.number_format = fmt


def _bar_chart(sheet, title: str, anchor: str) -> None:
    rows = sheet.max_row
    chart = BarChart()
    chart.type = "bar"
    chart.title = title
    chart.legend = None
    chart.height, chart.width = 9, 18
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=rows), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=rows))
    chart.y_axis.numFmt = MONEY
    chart.x_axis.scaling.orientation = "maxMin"  # biggest on top
    sheet.add_chart(chart, anchor)


def _line_chart(sheet, title: str, anchor: str) -> None:
    rows = sheet.max_row
    chart = LineChart()
    chart.title = title
    chart.legend = None
    chart.height, chart.width = 9, 22
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=rows), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=rows))
    chart.x_axis.number_format = "DD.MM"
    sheet.add_chart(chart, anchor)


__all__ = ["console_summary", "usage_rows", "write_chart", "write_excel", "write_powerbi"]
